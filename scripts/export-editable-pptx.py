"""Build native PowerPoint objects from the reviewed browser layout capture.
Requires python-pptx. Run from any directory; editable-layout.json retains geometry.
"""
import json, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

root=Path(__file__).resolve().parent.parent
deck=json.loads((root/'presentation/editable-layout.json').read_text())
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
prs.core_properties.title='KALINDO — Editable Destination Concept & Identity'
prs.core_properties.author='KALINDO Project'

def units(px):return Inches(max(0,px)/100)
def color(value):
    numbers=[float(n) for n in re.findall(r'[\d.]+',value)]
    return RGBColor(*[int(n) for n in numbers[:3]]), numbers[3] if len(numbers)>3 else 1

def fill(shape,value):
    rgb,alpha=color(value);shape.fill.solid();shape.fill.fore_color.rgb=rgb
    if alpha<1:
        el=shape.fill._xPr.solidFill.srgbClr
        a=OxmlElement('a:alpha');a.set('val',str(int(alpha*100000)));el.append(a)

def box(slide,x,y,w,h,c):
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,units(x),units(y),units(w),units(h))
    fill(shape,c);shape.line.fill.background();return shape

def text(slide,item):
    # Common fonts keep the file editable on typical PowerPoint installations.
    serif='Editorial' in item.get('font','')
    size=item['fontSize']*.72*(.94 if serif else .98)
    w=max(item['w']+20,item['fontSize']*1.1)
    h=item['h']+item['fontSize']*.45
    if item.get('container') and (serif or item['h']>item['lineHeight']*1.45):
        w=max(w,item['container']['w'])
    shape=slide.shapes.add_textbox(units(item['x']),units(item['y']-item['fontSize']*.10),units(w),units(h))
    shape.name='Editable text: '+item['text'][:65]
    tf=shape.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.TOP
    p=tf.paragraphs[0];p.text=item['text'];p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(item['lineHeight']*.72)
    p.alignment=PP_ALIGN.LEFT
    for run in p.runs:
        run.font.name='Georgia' if serif else 'Arial';run.font.size=Pt(size)
        run.font.bold=int(item.get('weight','400'))>=600
        run.font.color.rgb=color(item['color'])[0]
        if item.get('letterSpacing',0):run._r.get_or_add_rPr().set('spc',str(int(item['letterSpacing']*.72*100)))
    return shape

for index,data in enumerate(deck,1):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(data['background'])[0]
    for item in data['items']:
        kind=item['type']
        if kind=='rect':box(slide,item['x'],item['y'],item['w'],item['h'],item['color'])
        elif kind=='shade':
            shape=box(slide,item['x'],item['y'],item['w'],item['h'],'rgb(10,32,25)')
            sp=shape._element.spPr
            sp.remove(sp.solidFill)
            grad=OxmlElement('a:gradFill'); stops=OxmlElement('a:gsLst')
            for position,alpha in [('0','90000'),('100000','12000')]:
                stop=OxmlElement('a:gs');stop.set('pos',position)
                rgb=OxmlElement('a:srgbClr');rgb.set('val','0A2019')
                transparency=OxmlElement('a:alpha');transparency.set('val',alpha)
                rgb.append(transparency);stop.append(rgb);stops.append(stop)
            grad.append(stops);linear=OxmlElement('a:lin');linear.set('ang','0');linear.set('scaled','1');grad.append(linear)
            sp.append(grad)
        elif kind=='image':
            shape=slide.shapes.add_picture(str(root/item['src']),units(item['x']),units(item['y']),width=units(item['w']),height=units(item['h']))
            natural=item['naturalW']/item['naturalH'];target=item['w']/item['h']
            if natural>target:shape.crop_left=shape.crop_right=(1-target/natural)/2
            else:shape.crop_top=shape.crop_bottom=(1-natural/target)/2
            shape.name='Replaceable photo: '+Path(item['src']).name
        elif kind=='line':
            x,y,w,h=item['x'],item['y'],item['w'],item['h']
            if item['edge']=='Top':end=(x+w,y)
            elif item['edge']=='Bottom':y+=h;end=(x+w,y)
            elif item['edge']=='Left':end=(x,y+h)
            else:x+=w;end=(x,y+h)
            shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,units(x),units(y),units(end[0]),units(end[1]))
            shape.line.color.rgb=color(item['color'])[0];shape.line.width=Pt(item['width']*.72)
        elif kind=='text':text(slide,item)
        elif kind=='icon':
            # Editable, restrained line symbols replacing the browser SVG icons.
            x,y,w,h=item['x'],item['y'],item['w'],item['h'];path=item.get('key','')
            typ=MSO_SHAPE.OVAL if '14 14' in path or '6 6' in path else MSO_SHAPE.RECTANGLE
            shape=slide.shapes.add_shape(typ,units(x+7),units(y+7),units(w-14),units(h-14));shape.fill.background();shape.line.color.rgb=RGBColor.from_string('6D8767');shape.line.width=Pt(1)
    content=data.get('notes') or {}
    notes=[data['title'],'EDITABLE VERSION: text, photos, rules, diagrams and backgrounds are native PowerPoint objects. Click an element to edit, replace or delete it. All assets are embedded. Fonts use Georgia and Arial for portability.']
    for key in ('intro','text','second','quote','sub','note'):
        if content.get(key):notes.append(content[key])
    for item in content.get('items',[]):notes.append(' / '.join(item))
    notes.append('Source: KALINDO PROJECT.pdf. Concepts, programmes and naming recommendations are not approvals. See presentation/CONTENT-REVIEW.md.')
    slide.notes_slide.notes_text_frame.text='\n\n'.join(notes)
out=root/'KALINDO-Destination-Identity-EDITABLE.pptx';prs.save(out)
check=Presentation(out)
assert len(check.slides)==24
assert all(sum(s.has_text_frame for s in slide.shapes)>5 for slide in check.slides)
print(f'Created {out.name}: 24 slides, {sum(len(s.shapes) for s in check.slides)} native objects, {sum(s.has_text_frame for sl in check.slides for s in sl.shapes)} editable text/shape objects.')
