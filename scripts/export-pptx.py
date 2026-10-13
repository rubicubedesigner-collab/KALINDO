"""Export the reviewed slide renders to a faithful 16:9 PowerPoint.
Requires python-pptx and Pillow. Slide visuals are flattened; notes contain copy.
Run from the repository root after refreshing presentation/review screenshots.
"""
import json
import subprocess
from io import BytesIO
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parent.parent
result = subprocess.run(
    ['node', '--input-type=module', '-e',
     "import {slides} from './presentation/content.js'; console.log(JSON.stringify(slides));"],
    cwd=root, check=True, capture_output=True, text=True,
)
content = json.loads(result.stdout)
prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)
prs.core_properties.title = 'KALINDO — Destination Concept & Identity'
prs.core_properties.subject = 'Concept, Brand Positioning & Destination Naming'
prs.core_properties.author = 'KALINDO Project'
prs.core_properties.keywords = 'KALINDO, Batam, destination identity, brand concept'
for i, data in enumerate(content, 1):
    image_path = root / 'presentation' / 'review' / f'slide-{i:02}.png'
    with Image.open(image_path) as image:
        if image.width * 9 != image.height * 16:
            raise ValueError(f'Incorrect slide aspect ratio: {image_path}')
        buf = BytesIO()
        image.convert('RGB').save(buf, format='JPEG', quality=95, optimize=True)
        buf.seek(0)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = slide.shapes.add_picture(buf, 0, 0, width=prs.slide_width, height=prs.slide_height)
    picture.name = f'Slide {i:02} — {data["title"].replace(chr(10), " ")}'
    parts = [f'SLIDE {i:02} / {data["section"]}', data['title'],
             'Visual-fidelity version. The slide is a flattened image; edit source content/layout and regenerate to change it.']
    for key in ('eyebrow', 'intro', 'text', 'second', 'quote', 'sub'):
        if data.get(key):
            parts.append(data[key])
    for item in data.get('items', []):
        parts.append(' / '.join(str(value) for value in item))
    if data.get('note'):
        parts.append('REVIEW NOTE: ' + data['note'])
    if data['type'] == 'architecture':
        parts.append('Master Destination: [Selected Masterbrand]\nCategory Descriptor: Batam’s Tropical Lifestyle District\nSignature gathering & event experience: The Veranda\nProgramming: Veranda Sessions / Batam Table / The Weekend Edit / The Family Social\nExperience categories: Eat / Shop / Play / Move / Stay / Explore / Unwind')
    parts.append('Authority: KALINDO PROJECT.pdf. Supplied concept imagery is not verified as built. See presentation/CONTENT-REVIEW.md.')
    slide.notes_slide.notes_text_frame.text = '\n\n'.join(parts)
output = root / 'KALINDO-Destination-Identity.pptx'
prs.save(output)
check = Presentation(output)
assert len(check.slides) == 24
assert all(len(slide.shapes) == 1 for slide in check.slides)
assert all(slide.notes_slide.notes_text_frame.text for slide in check.slides)
print(f'Exported {output.name}: {len(check.slides)} slides, 16:9, with speaker notes.')
